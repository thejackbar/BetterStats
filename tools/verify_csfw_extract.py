"""Checks that the CSFW extractor is ONE tool for every club.

What differs between clubs is data - the record layout of the files, and the
club's own team and match-type names - so each check builds a small archive in a
layout and for a club the extractor was never written against, runs the shipped
code over it, and reads the answer back.

Every check reports rather than raising, so run it with a control (an older
extractor, `--module path/to/old.py`) and the missing parts are named instead of
ending the run.

    python3 tools/verify_csfw_extract.py [--module tools/csfw_extract.py]
"""
from __future__ import annotations

import collections
import csv
import datetime
import importlib.util
import json
import struct
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
args = sys.argv[1:]
MODULE = Path(args[args.index("--module") + 1]) if "--module" in args else HERE / "csfw_extract.py"

PASS = FAIL = 0


def check(label, got, want):
    global PASS, FAIL
    if got == want:
        PASS += 1
        print(f"PASS  {label}")
    else:
        FAIL += 1
        print(f"FAIL  {label}\n        got  {got!r}\n        want {want!r}")


def refuses(label, fn, needle):
    """`fn` must raise, and the message must say `needle`."""
    try:
        fn()
    except Exception as exc:                                   # noqa: BLE001
        check(label, needle if needle in str(exc) else f"{type(exc).__name__}: {exc}", needle)
        return
    check(label, "no error raised", needle)


def attempt(label, fn, want):
    """Run `fn`; a missing part or an exception is a FAIL that says what."""
    try:
        got = fn()
    except Exception as exc:                                   # noqa: BLE001
        got = f"{type(exc).__name__}: {exc}"
    check(label, got, want)


spec = importlib.util.spec_from_file_location("ext", MODULE)
m = importlib.util.module_from_spec(spec)
sys.modules["ext"] = m          # a dataclass looks its own module up here
spec.loader.exec_module(m)
EPOCH = datetime.date(1899, 12, 30)

# ---------------------------------------------------------------------------
# A layout the extractor has never seen: every table moved and resized, and the
# fields inside the records shifted. Built from the SAME field names the parser
# reads, so a parser that still had an offset hard-wired would read nonsense.
# ---------------------------------------------------------------------------
V2 = dict(
    name="synthetic-v2",
    player_base=400, player_rec=40, player_n=100,
    p_surname=4, p_initial=20, p_first=24,
    fixture_base=4_400, fixture_rec=70, fixture_n=60,
    f_date=4, f_team=30, f_name=34, f_name_len=30,
    match_base=8_600, match_rec=2_100, match_n=60,
    m_type=6, m_round=8, m_ground_idx=30, m_ground=40, m_ground_len=30,
    m_result=600, m_inn_us=610, m_inn_them=690, m_fow=780,
    block_base=830, block_rec=100, block_n=12,
    b_position=4, b_runs=6, b_sixes=10, b_fours=12, b_howout=18,
    b_overs=20, b_maidens=28, b_wickets=30, b_flags=36, b_conceded=38,
    b_wides=40, b_noballs=42, b_catches=44, b_wk_catches=48, b_stumpings=52,
    b_byes=56, b_votes=70, b_captain=60, b_keeper=62,
)
V2["file_len"] = V2["match_base"] + V2["match_n"] * V2["match_rec"]


def build_av(lay, club, year, fixtures, matches, players):
    """A .AV file in `lay`. fixtures: [(slot, date, team, name)] in file order."""
    buf = bytearray(lay.file_len)

    def i16(off, v):
        struct.pack_into("<h", buf, off, v)

    def cur(off, v):
        struct.pack_into("<q", buf, off, int(round(v * 10_000)))

    def txt(off, n, v):
        buf[off:off + n] = v.encode("latin-1").ljust(n)[:n]

    txt(lay.club_at, lay.club_len, club)
    txt(lay.year_at, lay.year_len, str(year))
    for i in range(lay.player_n):
        i16(lay.player_base + i * lay.player_rec, -1)
    for pid, (sur, first) in players.items():
        off = lay.player_base + pid * lay.player_rec
        i16(off, pid)
        txt(off + lay.p_surname, lay.p_surname_len, sur)
        txt(off + lay.p_first, lay.p_first_len, first)
    for i in range(lay.fixture_n):
        i16(lay.fixture_base + i * lay.fixture_rec, -1)
    for pos, (slot, date, team, name) in enumerate(fixtures):
        off = lay.fixture_base + pos * lay.fixture_rec
        i16(off, slot)
        struct.pack_into("<d", buf, off + lay.f_date, float((date - EPOCH).days))
        i16(off + lay.f_team, team)
        txt(off + lay.f_name, lay.f_name_len, name)
    for slot in range(lay.match_n):
        rec = lay.match_base + slot * lay.match_rec
        for inn in (lay.m_inn_us, lay.m_inn_them):
            for o in (lay.i_total, lay.i_wickets, lay.i_number, *lay.i_extras):
                i16(rec + inn + o, -1)
            cur(rec + inn + lay.i_overs, -1)
        for n in range(lay.m_fow_pairs):
            i16(rec + lay.m_fow + n * 4, -1)
        for b in range(lay.block_n):
            off = rec + lay.block_base + b * lay.block_rec
            i16(off + lay.b_player, -1)
            for o in (lay.b_runs, lay.b_sixes, lay.b_fours, lay.b_howout, lay.b_maidens,
                      lay.b_wickets, lay.b_flags, lay.b_conceded, lay.b_wides,
                      lay.b_noballs, lay.b_catches, lay.b_wk_catches,
                      lay.b_stumpings, lay.b_byes):
                i16(off + o, -1)
            i16(off + lay.b_position, lay.did_not_bat)
            cur(off + lay.b_overs, -1)
            cur(off + lay.b_votes, -1)
    for slot, mt in matches.items():
        rec = lay.match_base + slot * lay.match_rec
        buf[rec + lay.m_type] = mt["type"]
        buf[rec + lay.m_result] = mt["result"]
        txt(rec + lay.m_ground, lay.m_ground_len, mt.get("ground", "Home Oval"))
        for inn, key in ((lay.m_inn_us, "us"), (lay.m_inn_them, "them")):
            total, wkts = mt[key]
            i16(rec + inn + lay.i_total, total)
            i16(rec + inn + lay.i_wickets, wkts)
        for b, (pid, pos, runs, howout) in enumerate(mt["bat"]):
            off = rec + lay.block_base + b * lay.block_rec
            i16(off + lay.b_player, pid)
            i16(off + lay.b_position, pos)
            i16(off + lay.b_runs, runs)
            i16(off + lay.b_howout, howout)
        # Captain and wicketkeeper flags are Visual Basic booleans: -1 is True, 0 is False.
        for b, (cap, keep) in (mt.get("roles") or {}).items():
            off = rec + lay.block_base + b * lay.block_rec
            i16(off + lay.b_captain, -1 if cap else 0)
            i16(off + lay.b_keeper, -1 if keep else 0)
    return bytes(buf)


def squad(*pairs):
    return [(pid, pos, runs, 0 if runs % 2 else 11) for pid, pos, runs in pairs]


PLAYERS = {0: ("Smith", "Alan"), 1: ("Jones", "Ben"), 2: ("Patel", "Raj")}
FIXTURES = [  # file order is NOT slot order: the slot a fixture points at is the key
    (2, datetime.date(2025, 10, 25), 3, "Goodwood"),
    (0, datetime.date(2025, 10, 11), 1, "Fulham"),
    (1, datetime.date(2025, 10, 18), 4, "Glenunga"),
    (3, datetime.date(2025, 11, 1), 2, "Bye"),
]
MATCHES = {
    0: dict(type=4, result=0, us=(263, 7), them=(138, 10), ground="Devitt Oval",
            bat=squad((0, 0, 101), (1, 1, 41), (2, 2, 7))),
    1: dict(type=9, result=1, us=(90, 10), them=(91, 3), ground="Glenunga Oval",
            bat=squad((0, 0, 12), (1, 1, 30), (2, 2, 20))),
    2: dict(type=4, result=4, us=(-1, -1), them=(-1, -1), bat=squad((0, 0, 0))),
    3: dict(type=0, result=255, us=(-1, -1), them=(-1, -1), bat=[]),
}


def archive(folder, lay, files):
    """files: {name: (club, year)}"""
    for name, (club, year) in files.items():
        (folder / name).write_bytes(build_av(lay, club, year, FIXTURES, MATCHES, PLAYERS))


def run(*argv):
    p = subprocess.run([sys.executable, "-I", str(MODULE), *map(str, argv)],
                       capture_output=True, text=True)
    return p.returncode, p.stdout + p.stderr


def read_csv(path):
    return list(csv.DictReader(open(path, encoding="utf-8-sig")))


# ---------------------------------------------------------------------------
# 1. The layout is data
# ---------------------------------------------------------------------------
LAYOUT = getattr(m, "Layout", None)
if LAYOUT is None:
    check("the extractor describes a record layout as data", "no Layout in the module", "Layout")
    lay2 = None
else:
    lay2 = m.layout_from_dict(dict(V2))
    check("the extractor describes a record layout as data", lay2.fixture_rec, 70)

with tempfile.TemporaryDirectory() as td:
    td = Path(td)
    if lay2 is not None:
        (td / "DATA2025.AV").write_bytes(build_av(lay2, "Newtown CC", 2025, FIXTURES, MATCHES, PLAYERS))

        def unknown_msg():
            try:
                m.parse_file(td / "DATA2025.AV")
            except Exception as e:                              # noqa: BLE001
                return f"{type(e).__name__}: " + ("--analyse" if "--analyse" in str(e) else str(e))
            return "parsed without a layout"
        attempt("a file no layout describes is refused and points at --analyse",
                unknown_msg, "LayoutError: --analyse")

        # 2. described, then read by it
        lay_json = td / "v2.json"
        lay_json.write_text(json.dumps(V2))
        m.register_layouts(lay_json)
        s = m.parse_file(td / "DATA2025.AV")
        check("with its layout described, the file reads as that layout", s["layout"], "synthetic-v2")
        check("the club name comes from the layout's own header offset", s["club"], "Newtown CC")
        by_opp = {x["opponent"]: x for x in s["matches"]}
        check("fixtures are keyed by the slot they point at, not their position",
              sorted((x["opponent"], x["team"], x["date"].isoformat()) for x in s["matches"]),
              [("Bye", 2, "2025-11-01"), ("Fulham", 1, "2025-10-11"),
               ("Glenunga", 4, "2025-10-18"), ("Goodwood", 3, "2025-10-25")])
        check("the scorecard sits under its own fixture (Fulham 263/7, 101 by Smith)",
              (by_opp["Fulham"]["us"]["total"], by_opp["Fulham"]["blocks"][0]["runs"],
               s["players"][by_opp["Fulham"]["blocks"][0]["player_id"]]["surname"]),
              (263, 101, "Smith"))
        check("the club's result and match type are read at the layout's offsets",
              [(x["opponent"], x["result_code"], x["type_code"]) for x in s["matches"]],
              [("Fulham", 0, 4), ("Glenunga", 1, 9), ("Goodwood", 4, 4), ("Bye", 255, 0)])

        # 3. a layout is checked, not trusted
        refuses("an unknown layout field is refused, not ignored",
                lambda: m.layout_from_dict({**V2, "fixture_recc": 70}), "fixture_recc")
        refuses("a table running past the end of the file is refused",
                lambda: m.layout_from_dict({**V2, "match_n": 600}), "past the end")
        v1_size_garbage = bytearray(m.LAYOUT_V1.file_len)
        (td / "BAD2025.AV").write_bytes(bytes(v1_size_garbage))
        refuses("a file the right size for a layout but not shaped like it is refused",
                lambda: m.parse_file(td / "BAD2025.AV"), "not a year")

        # 4. the analyser finds the fixture table in a file nobody described
        many = [(i, datetime.date(2025, 10, 4) + datetime.timedelta(days=7 * i), (1, 3, 4)[i % 3],
                 f"Club {i}") for i in range(20)]
        info = m.probe_layout(build_av(lay2, "Newtown CC", 2025, many, {}, PLAYERS))
        fx = info["fixtures"] or {}
        check("--analyse finds the fixture table's record size in an unknown file",
              fx.get("record"), 70)
        cols = {c: v for c, v in fx.get("small_columns", [])}
        check("it flags the team column (few distinct values) by its distance from the date",
              {k for k in cols[V2["f_team"] - V2["f_date"]]} if V2["f_team"] - V2["f_date"] in cols else None,
              {1, 3, 4})
        check("it finds the opponent text beside it",
              (fx.get("text_at") or (None, None))[0], V2["f_name"] - V2["f_date"])

# ---------------------------------------------------------------------------
# 5. A club with no profile does not inherit Shoalwater Bay's names
# ---------------------------------------------------------------------------
if hasattr(m, "choose_profile"):
    prof = m.choose_profile("Newtown CC", None)
    m.activate(prof)
    check("another club's team number is unnamed, not Shoalwater's grade",
          m.grade_label(18), "Grade #18 (unmapped)")
    check("another club's match type is unnamed, not Shoalwater's 'Two day'",
          m.match_type_label(1), "type 1")
    check("an unnamed match type states no format (never a guess)",
          m.match_format_for(1, set(), 1, 50.0, 0.0), "")
    m.FINALS_TYPES = frozenset({9})
    m.FORMAT_BY_TYPE.clear()
    m.FORMAT_BY_TYPE.update({4: "One Day"})
    check("a final stays blank where only ONE format is named (it could be either)",
          m.match_format_for(9, set(), 1, 50.0, 0.0), "")
    m.FORMAT_BY_TYPE.update({1: "Two Day"})
    check("a final is measured once both formats are named",
          m.match_format_for(9, set(), 1, 50.0, 0.0), "One Day")
    sh = m.choose_profile("Shoalwater Bay Cricket Club", None)
    m.activate(sh)
    check("Shoalwater Bay's own profile is still chosen by its header name",
          (m.grade_label(18), m.match_type_label(4)), ("A grade", "One day"))
else:
    check("a club profile exists", "no choose_profile", "choose_profile")

# ---------------------------------------------------------------------------
# 6. End to end: two databases, an unprofiled club, a described layout
# ---------------------------------------------------------------------------
with tempfile.TemporaryDirectory() as td:
    td = Path(td)
    src, out = td / "src", td / "out"
    src.mkdir()
    if lay2 is not None:
        # DAT1 is the women's database THIS year; the name of a file says nothing
        (src / "DATA2025.AV").write_bytes(build_av(lay2, "Payneham", 2025, FIXTURES, MATCHES, PLAYERS))
        (src / "DAT12025.AV").write_bytes(build_av(lay2, "Payneham W", 2025, FIXTURES, MATCHES, PLAYERS))
    (td / "layout.json").write_text(json.dumps(V2))
    (td / "profile.json").write_text(json.dumps({
        "name": "payneham-test", "club": "payneham", "club_name": "Payneham Cricket Club",
        "team_names": {},
        "databases": {"Payneham": {"1": "1st XI", "3": "3rd XI"},
                      "Payneham W": {"1": "Women's 1st XI", "3": "Women's 3rd XI", "4": "Women's 4th XI"}},
        "match_types": {"4": "Limited Over", "9": "Final"}, "finals_types": [9],
        "format_by_type": {"4": "One Day"}, "byes": ["Bye"]}))
    code, text_out = run("--layout", td / "layout.json", "--analyse", src)
    check("--analyse runs on the folder", code, 0)
    check("--analyse groups by the database label in the header, with the file names beside it",
          all(w in text_out for w in ("'Payneham':", "'Payneham W':", "[DAT1]", "[DATA]")), True)
    check("the built-in Payneham profile is picked by the header, so team 1 of the men's database is named",
          "team   1 = 1st XI" in text_out, True)
    check("--analyse writes nothing", out.exists(), False)
    # a club nobody has a profile for: the bye is flagged, not silently imported
    (td / "other").mkdir()
    if lay2 is not None:
        (td / "other" / "DATA2025.AV").write_bytes(
            build_av(lay2, "Newtown CC", 2025, FIXTURES, MATCHES, PLAYERS))
    code, other_out = run("--layout", td / "layout.json", "--analyse", td / "other")
    check("--analyse hints at a bye with nothing played, for a club with no profile",
          "HINT: 1 fixture(s) against 'Bye'" in other_out, True)
    check("and a club with no profile is not given another club's names",
          "team   1 = Grade #1 (unmapped)" in other_out, True)

    code, text_out = run("--layout", td / "layout.json", td / "other", "-o", td / "out_other")
    check("extraction for a club with no profile still runs and says what is unnamed",
          (code, "NO NAME GIVEN FOR" in text_out), (0, True))
    code, text_out = run("--layout", td / "layout.json", "--profile", td / "profile.json",
                         src, "-o", out)
    check("extraction with a profile runs and drops the byes",
          (code, "2 bye fixtures dropped" in text_out), (0, True))
    check("ONE scorecard CSV for the whole folder, not one per file name",
          sorted(p.name for p in out.glob("manual_games_scorecards*.csv")),
          ["manual_games_scorecards.csv"])
    rows = read_csv(out / "manual_games_scorecards.csv") if code == 0 else []
    men = [r for r in rows if r["game_key"].startswith("data2025")]
    women = [r for r in rows if r["game_key"].startswith("dat12025")]
    check("each file's teams are named from ITS database (Fulham, team 1, in both)",
          ({r["grade_name"] for r in men if r["opposition"] == "Fulham"},
           {r["grade_name"] for r in women if r["opposition"] == "Fulham"}),
          ({"1st XI"}, {"Women's 1st XI"}))
    check("a team the profile did not name in that database is loud (Glenunga, team 4, men's)",
          ({r["grade_name"] for r in men if r["opposition"] == "Goodwood"},
           {r["grade_name"] for r in men if r["opposition"] == "Glenunga"},
           {r["grade_name"] for r in women if r["opposition"] == "Glenunga"}),
          ({"3rd XI"}, {"Grade #4 (unmapped)"}, {"Women's 4th XI"}))
    check("game keys carry the file name where several are in play, so they cannot collide",
          (men[0]["game_key"].split("-")[0] if men else None,
           women[0]["game_key"].split("-")[0] if women else None), ("data2025", "dat12025"))
    check("a bye is not a game",
          {r["opposition"] for r in rows}, {"Fulham", "Glenunga", "Goodwood"})
    check("the profile's club name is the winner when we win",
          {r["winning_team"] for r in men if r["opposition"] == "Fulham"}, {"Payneham Cricket Club"})
    check("the profile's match types decide format and finals; a final is blank with one format named",
          ({(r["opposition"], r["match_format"], r["is_final"]) for r in men}),
          {("Fulham", "One Day", ""), ("Glenunga", "", "true"), ("Goodwood", "One Day", "")})
    check("an abandoned match is blank, and still a match",
          {r["result"] for r in men if r["opposition"] == "Goodwood"}, {""})
    check("the club's result byte reaches the file (Fulham won, Glenunga lost)",
          ({r["result"] for r in men if r["opposition"] == "Fulham"},
           {r["result"] for r in men if r["opposition"] == "Glenunga"}),
          ({"WIN"}, {"LOSS"}))

# A single file name stem keeps the file names and game keys it always had.
with tempfile.TemporaryDirectory() as td:
    td = Path(td)
    src, out = td / "src", td / "out"
    src.mkdir()
    if lay2 is not None:
        archive(src, lay2, {"DATA2025.AV": ("Payneham CC", 2025)})
        (td / "layout.json").write_text(json.dumps(V2))
        code, _ = run("--layout", td / "layout.json", "--team-names", "1=1st XI,3=3rd XI,4=4th XI",
                      src, "-o", out)
        rows = read_csv(out / "manual_games_scorecards.csv") if code == 0 else []
        check("one file stem: the plain file name and the year-based game key",
              (code, rows[0]["game_key"] if rows else None), (0, "2025-001"))
        check("--team-names names the teams with no profile",
              sorted({r["grade_name"] for r in rows}), ["1st XI", "3rd XI", "4th XI"])

# Two layouts in one folder: each file is read with its own.
with tempfile.TemporaryDirectory() as td:
    td = Path(td)
    src = td / "src"
    src.mkdir()
    if lay2 is not None:
        (src / "DATA2024.AV").write_bytes(
            build_av(m.LAYOUT_V1, "Payneham CC", 2024, FIXTURES, MATCHES, PLAYERS))
        (src / "DATA2025.AV").write_bytes(build_av(lay2, "Payneham CC", 2025, FIXTURES, MATCHES, PLAYERS))
        (td / "layout.json").write_text(json.dumps(V2))
        code, text_out = run("--layout", td / "layout.json", "--analyse", src)
        check("a folder mixing two layouts reads each file with its own",
              ("csfw-av-v1: 1 file" in text_out, "synthetic-v2: 1 file" in text_out), (True, True))

# The two-space leg names Payneham's own files carry ("Reynella  1", "Reynella  2")
fx = {0: dict(date=1, team=3, name="Reynella  1"), 1: dict(date=1, team=3, name="Reynella  2"),
      2: dict(date=1, team=3, name="Reynella 12")}
attempt("two spaces in 'Reynella  1' / 'Reynella  2' do not split one match into two opponents",
        lambda: (m.split_leg(fx), sorted({(f["opponent"], f["leg"]) for f in fx.values()}))[1],
        [("Reynella", 1), ("Reynella", 2), ("Reynella 12", None)])

# ---------------------------------------------------------------------------
# 7. The club's answer sheet: written from the files, read back onto a profile
# ---------------------------------------------------------------------------
with tempfile.TemporaryDirectory() as td:
    td = Path(td)
    src = td / "src"
    src.mkdir()
    if lay2 is not None:
        (src / "DATA2025.AV").write_bytes(build_av(lay2, "Newtown CC", 2025, FIXTURES, MATCHES, PLAYERS))
    (td / "layout.json").write_text(json.dumps(V2))
    sheet = td / "sheet.csv"
    code, out_text = run("--layout", td / "layout.json", "--questions", sheet, src)
    check("--questions writes the sheet and stops (no extraction files)",
          (code, sheet.exists(), (td / "match_detail.xlsx").exists()), (0, True, False))
    rows = read_csv(sheet) if sheet.exists() else []
    check("the sheet's columns are the ones the covering letter names",
          list(rows[0].keys()) if rows else None,
          ["Section", "Ref", "Item", "Database", "Number", "Matches", "Years",
           "Example match", "What we have", "How to answer", "Your answer"])
    teams = [r for r in rows if r["Item"] == "Team"]
    check("one row per database and team number, with its matches and years",
          sorted((r["Database"], r["Number"], r["Matches"], r["Years"]) for r in teams),
          [("Newtown CC", "1", "1", "2025"), ("Newtown CC", "2", "1", "2025"),
           ("Newtown CC", "3", "1", "2025"), ("Newtown CC", "4", "1", "2025")])
    check("each row gives a match to recognise it by",
          sorted(r["Example match"] for r in teams),
          ["01 Nov 2025 v Bye", "11 Oct 2025 v Fulham", "18 Oct 2025 v Glenunga",
           "25 Oct 2025 v Goodwood"])
    (td / "byes.json").write_text(json.dumps({"name": "t", "club": "newtown", "byes": ["Bye"]}))
    sheet2 = td / "sheet2.csv"
    run("--layout", td / "layout.json", "--profile", td / "byes.json", "--questions", sheet2, src)
    rows2 = read_csv(sheet2) if sheet2.exists() else []
    check("with the bye in the profile its team is not asked about, and the bye is a question",
          (sorted(r["Number"] for r in rows2 if r["Item"] == "Team"),
           [(r["Item"], r["Matches"]) for r in rows2 if r["Item"] == "Bye"]),
          (["1", "3", "4"], [("Bye", "1")]))
    check("match types seen are asked, not types that were never used",
          sorted(r["Number"] for r in rows if r["Item"] == "Match type"), ["0", "4", "9"])
    check("every row says how to answer it",
          all(r["How to answer"] for r in rows), True)
    if not rows:
        check("the sheet has rows to answer", "none written", "rows")
    else:
        # the club fills it in
        for r in rows:
            if r["Item"] == "Team":
                r["Your answer"] = {"1": "Under 12 Blue", "2": "", "3": "Under 12 White", "4": ""}[r["Number"]]
            elif r["Item"] == "Match type":
                r["Your answer"] = {"0": "Friendly", "4": "Limited Over", "9": "Semi Final"}[r["Number"]]
        with sheet.open("w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
            w.writeheader()
            w.writerows(rows)
        out = td / "out"
        code, out_text = run("--layout", td / "layout.json", "--answers", sheet, src, "-o", out)
        games = read_csv(out / "manual_games_scorecards.csv") if code == 0 else []
        by = {g["opposition"]: g for g in games}
        check("answered teams are named, an unanswered one stays loudly unnamed",
              (code, by.get("Fulham", {}).get("grade_name"), by.get("Goodwood", {}).get("grade_name"),
               by.get("Glenunga", {}).get("grade_name")),
              (0, "Under 12 Blue", "Under 12 White", "Grade #4 (unmapped)"))
        check("an answered match type that says Final is a final, and its name is used",
              (by.get("Glenunga", {}).get("is_final"), by.get("Fulham", {}).get("is_final")), ("true", ""))
        # the same sheet read onto a profile directly, for what the CSV cannot hold
        if hasattr(m, "apply_answers"):
            extra = [dict(r, **{"Your answer": a}) for r, a in zip(
                [dict(r) for r in rows[:1]], ["x"])]
            more = rows + [
                {"Section": "C", "Ref": "C1", "Item": "Format", "Database": "", "Number": "4",
                 "Matches": "", "Years": "", "Example match": "", "What we have": "",
                 "How to answer": "", "Your answer": "One Day"},
                {"Section": "C", "Ref": "C2", "Item": "Absent batter", "Database": "", "Number": "",
                 "Matches": "", "Years": "", "Example match": "", "What we have": "",
                 "How to answer": "", "Your answer": "Did not bat"},
                {"Section": "C", "Ref": "C3", "Item": "Result code", "Database": "", "Number": "6",
                 "Matches": "", "Years": "", "Example match": "", "What we have": "",
                 "How to answer": "", "Your answer": "No result"},
                {"Section": "C", "Ref": "C4", "Item": "Bye", "Database": "", "Number": "",
                 "Matches": "", "Years": "", "Example match": "", "What we have": "",
                 "How to answer": "", "Your answer": "No"}]
            full = td / "full.csv"
            with full.open("w", newline="", encoding="utf-8-sig") as fh:
                w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
                w.writeheader()
                w.writerows(more)
            base = m.Profile("t", byes=("bye",))
            prof2, absent, notes = m.apply_answers(base, full)
            check("a Format answer sets the format, an Absent answer sets the reading, a No keeps the byes",
                  (prof2.format_by_type, absent, prof2.byes), ({4: "One Day"}, "did-not-bat", ()))
            check("what no profile field can hold is handed back for a person, not guessed",
                  [n for n in notes if "Result code 6" in n] != [], True)
            check("a blank answer changes nothing; the base profile is not edited in place",
                  (base.byes, base.format_by_type, base.databases), (("bye",), {}, {}))
        else:
            check("the answer sheet can be read back", "no apply_answers", "apply_answers")


# ---------------------------------------------------------------------------
# 8. What a club answers: result codes, absent batters, layered team names, and
#    a workbook with CSFW's own Teams and Match Types tabs
# ---------------------------------------------------------------------------
with tempfile.TemporaryDirectory() as td:
    td = Path(td)
    src = td / "src"
    src.mkdir()
    fx = [(0, datetime.date(2025, 10, 11), 1, "Fulham"), (1, datetime.date(2025, 10, 18), 1, "Glenunga"),
          (2, datetime.date(2025, 10, 25), 1, "Goodwood"), (3, datetime.date(2025, 11, 1), 1, "Unley")]
    mt = {
        0: dict(type=4, result=5, us=(100, 10), them=(90, 10), bat=[(0, 0, 10, 11)]),
        1: dict(type=4, result=6, us=(100, 10), them=(100, 10), bat=[(0, 0, 10, 11)]),
        2: dict(type=4, result=7, us=(120, 10), them=(100, 10), bat=[(0, 0, 10, 11)]),
        3: dict(type=4, result=0, us=(120, 10), them=(100, 10), bat=[(0, 0, 0, 29), (1, 1, 12, 11)]),
    }
    if lay2 is not None:
        (src / "DATA2025.AV").write_bytes(build_av(lay2, "Newtown", 2025, fx, mt, PLAYERS))
        (src / "DAT12025.AV").write_bytes(build_av(lay2, "Newtown W", 2025, fx[:1], {0: mt[0]}, PLAYERS))
    (td / "layout.json").write_text(json.dumps(V2))
    base = {"name": "t", "club": "newtown", "team_names": {"1": "1st XI"}, "byes": ["Bye"],
            "databases": {"Newtown W": {"1": "Women's 1st XI"}}}
    (td / "plain.json").write_text(json.dumps(base))
    (td / "answered.json").write_text(json.dumps({
        **base, "absent": "did-not-bat",
        "result_labels": {"5": "Cancelled", "6": "Tied on first innings", "7": "Winning draw"},
        "import_results": {"5": "", "6": "DRAW", "7": "DRAW"}}))

    def extract(profile):
        out = td / ("out_" + profile.stem)
        code, text_out = run("--layout", td / "layout.json", "--profile", profile, src, "-o", out)
        return code, (read_csv(out / "manual_games_scorecards.csv") if code == 0 else [])

    code, rows = extract(td / "plain.json")
    by = {(r["game_key"][:4], r["opposition"]): r for r in rows}
    check("a result code nobody has named loads with no result (never a guessed one)",
          [by.get(("data", o), {}).get("result") for o in ("Fulham", "Glenunga", "Goodwood")], ["", "", ""])
    check("layered names: the global list names every database, a database's own entry wins",
          (by.get(("data", "Fulham"), {}).get("grade_name"), by.get(("dat1", "Fulham"), {}).get("grade_name")),
          ("1st XI", "Women's 1st XI"))
    check("an absent batter reads as 'absent out' until the club says otherwise",
          [r["did_not_bat"] for r in rows if r["opposition"] == "Unley" and r["player_name"] == "Smith, Alan"],
          ["false"])
    code, rows = extract(td / "answered.json")
    by = {(r["game_key"][:4], r["opposition"]): r for r in rows}
    check("the club's answers set what a result means: cancelled blank, tied-on-first-innings and winning draw are draws",
          [by.get(("data", o), {}).get("result") for o in ("Fulham", "Glenunga", "Goodwood", "Unley")],
          ["", "DRAW", "DRAW", "WIN"])
    check("the club's answer that an absent batter did not bat is applied",
          [r["did_not_bat"] for r in rows if r["opposition"] == "Unley" and r["player_name"] == "Smith, Alan"],
          ["true"])
    if hasattr(m, "apply_answers"):
        # the follow-up sheet asks only what is still open
        code, _ = run("--layout", td / "layout.json", "--profile", td / "answered.json",
                      "--questions", td / "open.csv", "--open-only", src)
        open_rows = read_csv(td / "open.csv") if (td / "open.csv").exists() else []
        check("the follow-up sheet leaves out teams and results the profile has, and keeps a type still unnamed",
              sorted({r["Item"] for r in open_rows}), ["Match type"])
        # a workbook with CSFW's own tabs
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Answer sheet"
        ws.append(["Section", "Ref", "Item", "Database", "Number", "Your answer"])
        ws.append(["A", "A1", "Team", "Newtown", 1, "1st xi"])
        ws.append(["A", "A2", "Team", "Newtown", 2, "Seconds"])
        ws.append(["A", "A3", "Team", "Newtown", 3, "Under 11"])
        ws.append(["A", "A4", "Team", "Newtown W", 9, "Women's 9th"])
        ws.append(["A", "A5", "Team", "Newtown", 7, ""])
        t = wb.create_sheet("Teams")
        t.append(["ID", "Team"])
        for i, n in ((1, "1st XI"), (2, "2nd Xi"), (3, "U 11 Blue")):
            t.append([i, n])
        mtab = wb.create_sheet("Match Types")
        mtab.append(["ID", "Match Type"])
        for i, n in ((0, "Friendly"), (3, "Cup Final"), (4, "Limited Over")):
            mtab.append([i, n])
        wb.save(td / "back.xlsx")
        prof, absent, notes = m.apply_answers(m.Profile("t"), td / "back.xlsx")
        check("a Teams tab names the teams for every database; the club's own spelling of the same name is kept",
              (prof.team_names[1], prof.team_names[2], prof.team_names[3]),
              ("1st xi", "2nd Xi", "U 11 Blue"))
        check("a row answer that disagrees with the tab is reported, and the tab is used",
              ([n for n in notes if "team 3" in n] != [], prof.team_names[3]), (True, "U 11 Blue"))
        check("a row for a number the tab lacks is kept against its own database",
              prof.databases.get("Newtown W"), {9: "Women's 9th"})
        check("a Match Types tab names the types, and a name with Final in it is a final",
              (prof.match_types, sorted(prof.finals_types)),
              ({0: "Friendly", 3: "Cup Final", 4: "Limited Over"}, [3]))

# ---------------------------------------------------------------------------
# 9. A match type that can be run as either format is settled game by game
# ---------------------------------------------------------------------------
with tempfile.TemporaryDirectory() as td:
    td = Path(td)
    src = td / "src"
    src.mkdir()
    d = datetime.date(2025, 11, 1)
    fx = [(0, d, 1, "Alpha 1"), (1, d, 1, "Alpha 2"),          # type 8, a second innings recorded
          (2, d, 2, "Beta"),                                    # type 8, one innings
          (3, d, 3, "Gamma 1"), (4, d, 3, "Gamma 2"),          # type 13, a second innings recorded
          (5, d, 4, "Delta"),                                   # type 13, one innings
          (6, d, 4, "Washout")]                                 # no type, no score
    bat = [(0, 0, 10, 11)]
    mt = {0: dict(type=8, result=0, us=(100, 10), them=(90, 10), bat=bat),
          1: dict(type=8, result=0, us=(60, 5), them=(40, 10), bat=bat),
          2: dict(type=8, result=0, us=(100, 10), them=(90, 10), bat=bat),
          3: dict(type=13, result=0, us=(100, 10), them=(90, 10), bat=bat),
          4: dict(type=13, result=0, us=(60, 5), them=(40, 10), bat=bat),
          5: dict(type=13, result=0, us=(100, 10), them=(90, 10), bat=bat),
          6: dict(type=200, result=4, us=(-1, -1), them=(-1, -1), bat=[])}
    if lay2 is not None:
        (src / "DATA2025.AV").write_bytes(build_av(lay2, "Newtown", 2025, fx, mt, PLAYERS))
    (td / "layout.json").write_text(json.dumps(V2))
    base = {"name": "t", "club": "newtown", "team_names": {"1": "A", "2": "B", "3": "C", "4": "D"},
            "match_types": {"8": "Junior", "13": "Old comp"},
            "format_by_type": {"8": "One Day", "13": "Two Day"}}
    (td / "plain.json").write_text(json.dumps(base))
    (td / "mixed.json").write_text(json.dumps({**base, "mixed_format_types": [8, 13]}))

    def formats(profile):
        out = td / ("o_" + profile.stem)
        code, _ = run("--layout", td / "layout.json", "--profile", profile, src, "-o", out)
        rows = read_csv(out / "manual_games_scorecards.csv") if code == 0 else []
        return {r["opposition"]: r["match_format"] for r in rows}

    check("without it, every game of a type takes the type's one format (the old behaviour)",
          formats(td / "plain.json"),
          {"Alpha": "One Day", "Beta": "One Day", "Gamma": "Two Day", "Delta": "Two Day", "Washout": ""})
    check("for a mixed type a recorded second innings makes the game Two Day, and the rest keep the club's "
          "One Day, or stay blank where the club's answer was Two Day (nothing proves a two-day game ran one day)",
          formats(td / "mixed.json"),
          {"Alpha": "Two Day", "Beta": "One Day", "Gamma": "Two Day", "Delta": "", "Washout": ""})

    # the games left with no format are listed for the club, and its answers come back per game
    out = td / "o_mixed"
    review = read_csv(out / "games_without_format.csv") if (out / "games_without_format.csv").exists() else []
    check("the games with no format are listed for review (and an abandoned no-type game is not)",
          [(r["Opposition"], r["Match type"], r["Second innings on the card"]) for r in review],
          [("Delta", "Old comp", "No")])
    check("each listed game says why it is blank and has a game key to match its answer to",
          (bool(review and review[0]["Why blank"]), review[0]["Game key"] if review else None),
          (True, "2025-004"))
    if review:
        review[0]["Your answer"] = "One Day"
        filled = td / "review_filled.csv"
        with filled.open("w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=list(review[0].keys()))
            w.writeheader()
            w.writerows(review)
        out2 = td / "o_answered"
        run("--layout", td / "layout.json", "--profile", td / "mixed.json", "--answers", filled,
            src, "-o", out2)
        got = {r["opposition"]: r["match_format"] for r in read_csv(out2 / "manual_games_scorecards.csv")}
        check("the club's answer for one game sets that game's format, and nothing else moves",
              got, {"Alpha": "Two Day", "Beta": "One Day", "Gamma": "Two Day", "Delta": "One Day", "Washout": ""})
        check("once every listed game is answered the review sheet is not written",
              (out2 / "games_without_format.csv").exists(), False)
        # Payneham typed over the instruction column instead of "Your answer"
        review[0]["Your answer"] = ""
        review[0]["How to answer"] = "Two Day"
        wrong = td / "review_wrong_column.csv"
        with wrong.open("w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=list(review[0].keys()))
            w.writeheader()
            w.writerows(review)
        prof3, _, notes3 = m.apply_answers(m.Profile("t"), wrong) if hasattr(m, "apply_answers") else (None, None, [])
        check("an answer typed over the instruction column is read, and the club is told",
              (prof3.game_formats if prof3 else None, [n for n in notes3 if "How to answer" in n] != []),
              ({"2025-004": "Two Day"}, True))
        review[0]["How to answer"] = "Type Two Day or One Day"
        plain = td / "review_unanswered.csv"
        with plain.open("w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=list(review[0].keys()))
            w.writeheader()
            w.writerows(review)
        prof4, _, _ = m.apply_answers(m.Profile("t"), plain) if hasattr(m, "apply_answers") else (None, None, [])
        check("the instruction text itself is never taken for an answer",
              prof4.game_formats if prof4 else None, {})

# ---------------------------------------------------------------------------
# 10. One person listed twice in one innings
# ---------------------------------------------------------------------------
# CSFW can hold the same player id in two block slots of one match. The
# importer refuses a second row for the same player and innings, and drops the
# whole match for it (Payneham's first load lost 17 matches that way).
with tempfile.TemporaryDirectory() as td:
    td = Path(td)
    twice = {
        0: dict(type=4, result=0, us=(30, 3), them=(20, 10), ground="Home Oval",
                bat=[(0, 0, 10, 1), (1, 1, 15, 11), (0, 4, 3, 11), (2, 2, 2, 1)]),
    }
    (td / "DATA2025.AV").write_bytes(build_av(
        m.layout_from_dict(dict(V2)), "Dup CC", 2025,
        [(0, datetime.date(2025, 10, 11), 1, "Fulham")], twice, PLAYERS))
    seasons = [m.parse_file(td / "DATA2025.AV")]
    leg = seasons[0]["matches"][0]
    mine = [b for b in leg["blocks"] if b["player_id"] == 0]
    check("a person listed twice in a leg is one block, with the figures added",
          (len(mine), mine[0]["runs"], mine[0]["position"], mine[0]["not_out"]),
          (1, 13, 1, False))
    check("what was folded is recorded on the match for the club to check",
          [(x["player_id"], x["entries"], x["runs"]) for x in leg.get("listed_twice", [])],
          [(0, 2, [10, 3])])
    check("the others in the leg are untouched",
          sorted((b["player_id"], b["runs"]) for b in leg["blocks"]),
          [(0, 13), (1, 15), (2, 2)])
    rows = m.build_game_rows(seasons, "Dup CC")
    merge = getattr(m, "merge_duplicate_blocks", None)
    per = collections.Counter((r["game_key"], r["innings_number"], r["player_name"])
                              for r in rows if r["player_name"])
    check("no player has two rows in one innings (the importer would drop the match)",
          max(per.values()), 1)
    check("the folded row carries the added runs",
          [r["batting_runs"] for r in rows if r["player_name"] == "Smith, Alan"], [13])
    check("and the review list names the person",
          [(r["Player"], r["Entries"]) for r in getattr(m, "listed_twice_rows", lambda _: [])(seasons)],
          [("Smith, Alan", 2)])
    # a batting entry and a separate bowling entry fold into one with both
    blocks = [
        dict(slot=0, player_id=5, position=1, batted=True, runs=20, fours=2, sixes=0,
             not_out=False, dismissal_code=1, bowled=False, overs=None, balls=0,
             maidens=None, wickets=None, conceded=None, wides=None, no_balls=None,
             catches=1, wk_catches=None, stumpings=None, byes=None, votes=None, flags=None),
        dict(slot=7, player_id=5, position=None, batted=False, runs=None, fours=None,
             sixes=None, not_out=False, dismissal_code=None, bowled=True, overs=4.3,
             balls=27, maidens=1, wickets=2, conceded=18, wides=1, no_balls=None,
             catches=None, wk_catches=None, stumpings=None, byes=None, votes=None, flags=None),
    ]
    merged, info = merge(blocks) if merge else (blocks, [])
    one = merged[0]
    check("a batting entry and a bowling entry become one row with both",
          (len(merged), one["runs"], one["overs"], one["wickets"], one["catches"],
           one["position"], one["batted"], one["bowled"]),
          (1, 20, 4.3, 2, 1, 1, True, True))
    check("a leg with nobody listed twice comes back as it went in",
          merge(blocks[:1]) if merge else None, (blocks[:1], []))

# ---------------------------------------------------------------------------
# 11. Captain and wicketkeeper (player block offsets 4 and 6, Visual Basic booleans)
# ---------------------------------------------------------------------------
# Read at the layout's own offsets (the V2 layout puts them elsewhere, so a parser that still had 4 and 6
# hard-wired would read the batting position as a captain), carried to the scorecard CSV, and written once per
# person per match to the roles file the apply script reads.
with tempfile.TemporaryDirectory() as td:
    td = Path(td)
    flagged = {
        # Jones captains, Smith keeps; Smith is listed twice and only one of his entries carries the flag
        0: dict(type=4, result=0, us=(30, 3), them=(20, 10), ground="Home Oval",
                bat=[(0, 0, 10, 1), (1, 1, 15, 11), (2, 2, 2, 1), (0, 4, 3, 11)],
                roles={1: (True, False), 0: (False, False), 3: (False, True)}),
        # two keepers (a change of gloves), no captain
        1: dict(type=4, result=1, us=(50, 5), them=(60, 2), ground="Home Oval",
                bat=[(0, 0, 20, 1), (1, 1, 20, 11), (2, 2, 10, 1)],
                roles={0: (False, True), 2: (False, True)}),
        # nobody flagged
        2: dict(type=4, result=0, us=(70, 5), them=(60, 2), ground="Home Oval",
                bat=[(0, 0, 30, 1), (1, 1, 40, 11)]),
    }
    (td / "DATA2025.AV").write_bytes(build_av(
        m.layout_from_dict(dict(V2)), "Role CC", 2025,
        [(0, datetime.date(2025, 10, 11), 1, "Fulham"), (1, datetime.date(2025, 10, 18), 1, "Glenunga"),
         (2, datetime.date(2025, 10, 25), 1, "Goodwood")], flagged, PLAYERS))
    seasons = [m.parse_file(td / "DATA2025.AV")]
    legs = {x["opponent"]: x for x in seasons[0]["matches"]}
    check("the flags are read at the layout's offsets: Jones captain, Smith (listed twice, one entry flagged) the keeper",
          sorted((seasons[0]["players"][b["player_id"]]["surname"], b["captain"], b["keeper"])
                 for b in legs["Fulham"]["blocks"]),
          [("Jones", True, False), ("Patel", False, False), ("Smith", False, True)])
    rows = m.build_game_rows(seasons, "Role CC")
    cols = m.GAME_CSV_COLUMNS
    check("the scorecard CSV has the two columns", ("is_captain" in cols, "is_wicketkeeper" in cols), (True, True))
    by = {(r["opposition"], r["player_name"]): r for r in rows if r["player_name"]}
    check("Jones is the captain on the row, true or blank",
          (by[("Fulham", "Jones, Ben")]["is_captain"], by[("Fulham", "Jones, Ben")]["is_wicketkeeper"]), ("true", ""))
    check("Smith listed twice is one person, and the keeper flag on either entry makes him the keeper",
          (by[("Fulham", "Smith, Alan")]["is_wicketkeeper"], by[("Fulham", "Smith, Alan")]["is_captain"]), ("true", ""))
    check("a match with nobody flagged leaves both columns blank",
          {(r["is_captain"], r["is_wicketkeeper"]) for r in rows if r["opposition"] == "Goodwood"}, {("", "")})
    roles = m.game_role_rows(rows)
    check("the roles file has one row per person holding a role, none for the rest",
          sorted((r["opposition"], r["player_name"], r["is_captain"], r["is_wicketkeeper"]) for r in roles),
          [("Fulham", "Jones, Ben", "true", ""), ("Fulham", "Smith, Alan", "", "true"),
           ("Glenunga", "Patel, Raj", "", "true"), ("Glenunga", "Smith, Alan", "", "true")])
    check("and carries what finds the match again (season, grade, date, opposition)",
          sorted(set(roles[0]) ^ set(m.ROLE_COLUMNS)), [])
    rep = m.role_report(rows)
    check("the report counts matches, captains, keepers and two-keeper matches",
          (rep["matches"], rep["with_captain"], rep["with_keeper"], rep["two_keepers"], len(rep["two_captains"])),
          (3, 1, 2, 1, 0))
    flagged[2]["roles"] = {0: (True, False), 1: (True, False)}
    (td / "DATA2025.AV").write_bytes(build_av(
        m.layout_from_dict(dict(V2)), "Role CC", 2025,
        [(0, datetime.date(2025, 10, 11), 1, "Fulham"), (1, datetime.date(2025, 10, 18), 1, "Glenunga"),
         (2, datetime.date(2025, 10, 25), 1, "Goodwood")], flagged, PLAYERS))
    two = m.role_report(m.build_game_rows([m.parse_file(td / "DATA2025.AV")], "Role CC"))
    check("a match flagging two captains is listed for the club to check",
          [(x["row"]["opposition"], x["captains"]) for x in two["two_captains"]],
          [("Goodwood", ["Smith, Alan", "Jones, Ben"])])

print(f"\n{PASS} passed, {FAIL} failed")
sys.exit(1 if FAIL else 0)
